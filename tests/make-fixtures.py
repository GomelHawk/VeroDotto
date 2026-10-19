# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "pypandoc_binary==1.17",
#   "odfpy==1.4.1",
#   "openpyxl==3.1.5",
#   "python-pptx==1.0.2",
#   "fpdf2==2.8.4",
# ]
# ///
"""Regenerates tests/fixtures: one tiny file per format, plus what its markdown must and
must not contain. Run by hand, outputs committed:

    uv run --script tests/make-fixtures.py

Everything is generated here, so there is no licensing question about any fixture.
"""

from __future__ import annotations

import struct
import zlib
from email.message import EmailMessage
from pathlib import Path

import pypandoc
from fpdf import FPDF
from odf import draw, presentation, style, table, text
from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent / "fixtures"
REJECT_COMMON = ["<w:", "<a:", "<office:", "\r"]


def png(width: int, height: int, rgb: tuple[int, int, int]) -> bytes:
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def expect(name: str, must: list[str], reject: list[str] | None = None) -> None:
    (OUT / f"{name}.expect").write_text("\n".join(must) + "\n", encoding="utf-8")
    (OUT / f"{name}.reject").write_text("\n".join(REJECT_COMMON + (reject or [])) + "\n", encoding="utf-8")


SPEC_MD = """# Product specification

## Goals

- Read documents as markdown
- Spend fewer tokens

| Format | Converter |
| --- | --- |
| docx | markitdown |
| odt | pandoc |

![A red square](red.png)
"""


def main() -> None:
    OUT.mkdir(exist_ok=True)
    work = OUT / ".work"
    work.mkdir(exist_ok=True)
    (work / "red.png").write_bytes(png(8, 8, (220, 30, 30)))

    def from_md(md: str, fmt: str, dest: str) -> None:
        pypandoc.convert_text(md, fmt, format="md", outputfile=str(OUT / dest), extra_args=[f"--resource-path={work}"])

    from_md(SPEC_MD, "docx", "spec.docx")
    expect("spec.docx", ["Product specification", "Spend fewer tokens", "markitdown"])

    from_md("# Отчёт о работе\n\nЭто проверка оценки токенов для кириллицы.\n", "docx", "cyrillic.docx")
    expect("cyrillic.docx", ["Отчёт о работе", "кириллицы"])

    from_md(SPEC_MD, "odt", "doc.odt")
    expect("doc.odt", ["Product specification", "Spend fewer tokens"])

    # --standalone: without it pandoc writes an RTF fragment with no {\rtf1 header
    pypandoc.convert_text(
        "# Rich text\n\nA paragraph in **RTF**.\n", "rtf", format="md", outputfile=str(OUT / "note.rtf"), extra_args=["--standalone"]
    )
    expect("note.rtf", ["Rich text", "paragraph in"], ["\\rtf1"])

    (OUT / "page.html").write_text(
        "<html><head><style>.x{color:red}</style><script>var tracking = 1;</script></head><body>"
        "<nav><a href='/'>Home</a> | <a href='/about'>About</a></nav>"
        "<h1>Release notes</h1><p>Version <b>1.0</b> reads documents.</p>"
        "<img src='shot.png' alt='screenshot'></body></html>",
        encoding="utf-8",
    )
    expect("page.html", ["Release notes", "reads documents"], ["tracking", "color:red"])

    pypandoc.convert_text(
        "---\ntitle: A small book\n---\n\n# Chapter one\n\nIt was a quiet morning.\n\n# Chapter two\n\nThe end.\n",
        "epub",
        format="md",
        outputfile=str(OUT / "book.epub"),
    )
    expect("book.epub", ["Chapter one", "quiet morning", "Chapter two"])

    deck = Presentation()
    for title, body, notes in [
        ("Quarterly review", "Revenue grew", "Mention the new customers"),
        ("Next steps", "Hire two engineers", "Ask about budget"),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        slide.notes_slide.notes_text_frame.text = notes
    deck.slides[0].shapes.add_picture(str(work / "red.png"), Inches(1), Inches(5), Inches(1), Inches(1))
    deck.save(OUT / "deck.pptx")
    expect("deck.pptx", ["Quarterly review", "Next steps", "Hire two engineers", "Mention the new customers"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Laptops", 4200])
    ws.append(["Licences", 900])
    ws.append(["Total", "=SUM(B2:B3)"])
    people = wb.create_sheet("People")
    people.append(["Name", "Role"])
    people.append(["Ada", "Engineer"])
    wb.save(OUT / "book.xlsx")
    # openpyxl saves no formula results: the total must come through as its formula
    expect("book.xlsx", ["Budget", "People", "Laptops", "Engineer", "| Total | =SUM(B2:B3) |"])

    big = Workbook()
    ws = big.active
    ws.title = "Readings"
    ws.append(["Row", "Sensor", "Value"])
    for i in range(1, 10001):
        ws.append([i, f"S{i % 7}", i * 0.5])
    big.save(OUT / "big.xlsx")
    expect("big.xlsx", ["(summary)", "Readings", "10001 rows", "Sensor"], ["| 5000 |"])

    ods = OpenDocumentSpreadsheet()
    sheet = table.Table(name="Inventory")
    for cells in [["Part", "Count"], ["Bolt", "12"], ["Nut", "40"]]:
        row = table.TableRow()
        for value in cells:
            cell = table.TableCell(valuetype="string")
            cell.addElement(text.P(text=value))
            row.addElement(cell)
        sheet.addElement(row)
    ods.spreadsheet.addElement(sheet)
    ods.save(str(OUT / "sheet.ods"))
    expect("sheet.ods", ["Inventory", "Bolt", "Nut"])

    odp = OpenDocumentPresentation()
    layout = style.PageLayout(name="L")
    odp.automaticstyles.addElement(layout)
    master = style.MasterPage(name="M", pagelayoutname=layout)
    odp.masterstyles.addElement(master)
    for title, notes in [("Roadmap", "Keep it short"), ("Risks", "")]:
        page = draw.Page(masterpagename=master, name=title)
        frame = draw.Frame(width="20cm", height="3cm", x="1cm", y="1cm")
        box = draw.TextBox()
        box.addElement(text.P(text=f"{title} slide text"))
        frame.addElement(box)
        page.addElement(frame)
        if notes:
            note = presentation.Notes()
            nframe = draw.Frame(width="20cm", height="3cm", x="1cm", y="10cm")
            nbox = draw.TextBox()
            nbox.addElement(text.P(text=notes))
            nframe.addElement(nbox)
            note.addElement(nframe)
            page.addElement(note)
        odp.presentation.addElement(page)
    odp.save(str(OUT / "slides.odp"))
    expect("slides.odp", ["Roadmap slide text", "Risks slide text", "Keep it short"])

    pdf = FPDF()
    for n in range(1, 3):
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(text=f"Annual report, page {n}")
        pdf.ln(10)
        for i in range(12):
            pdf.multi_cell(0, 6, text=f"Paragraph {i} on page {n}: the numbers went up and to the right this year.", new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(OUT / "text.pdf"))
    expect("text.pdf", ["Annual report, page 1", "page 2"])

    grid = FPDF()
    grid.add_page()
    grid.set_font("Helvetica", size=11)
    grid.multi_cell(0, 6, text="Sales by region, with the table below.", new_x="LMARGIN", new_y="NEXT")
    with grid.table() as tbl:
        for cells in [("Region", "Q1", "Q2"), ("North", "120", "135"), ("South", "98", "101")]:
            row = tbl.row()
            for value in cells:
                row.cell(value)
    for i in range(10):
        grid.multi_cell(0, 6, text=f"Commentary line {i}: totals are rounded to whole units.", new_x="LMARGIN", new_y="NEXT")
    grid.output(str(OUT / "table.pdf"))
    expect("table.pdf", ["Sales by region", "| Region | Q1 | Q2 |", "| North | 120 | 135 |", "Commentary line 9"])

    scan = FPDF()
    scan.add_page()
    scan.image(str(work / "red.png"), x=10, y=10, w=180)
    scan.output(str(OUT / "scan.pdf"))
    expect("scan.pdf", [])

    mail = EmailMessage()
    mail["From"] = "Ada <ada@example.com>"
    mail["To"] = "team@example.com"
    mail["Subject"] = "Launch checklist"
    mail.set_content("Hi all,\n\nThe launch is on Friday.\n")
    mail.add_alternative("<p>Hi all,</p><p>The launch is on <b>Friday</b>.</p>", subtype="html")
    mail.add_attachment(b"col\n1\n", maintype="text", subtype="csv", filename="plan.csv")
    (OUT / "mail.eml").write_bytes(bytes(mail))
    expect("mail.eml", ["Launch checklist", "launch is on", "plan.csv"])

    for leftover in work.iterdir():
        leftover.unlink()
    work.rmdir()


if __name__ == "__main__":
    main()
