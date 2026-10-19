# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "mammoth==1.13.0",
#   "markdownify==1.2.3",
#   "beautifulsoup4==4.15.0",
#   "python-pptx==1.0.2",
#   "openpyxl==3.1.5",
#   "xlrd==2.0.2",
#   "odfpy==1.4.1",
#   "pdfplumber==0.11.10",
#   "striprtf==0.0.33",
#   "olefile==0.47",
# ]
# ///
"""VeroDotto's converters: one document in, doc.md and meta.json out.

    convert.py convert SRC OUTDIR [--max-sheet-cells N] [--max-sheet-rows N]
                                  [--max-tokens N] [--timeout S]
    convert.py --selftest     convert one built-in sample per format, print versions

Run by verodotto.sh through `uv run --script`, which installs the pinned dependencies
above into VeroDotto's private cache. A lean stack by design: the libraries markitdown
itself runs underneath (mammoth, python-pptx, openpyxl, pdfplumber, ...), with no pandoc,
pandas or onnxruntime - about a third of the size. The token counts are estimates
(README: "How the numbers are estimated") and say so wherever they are shown.
"""

from __future__ import annotations

import argparse
import datetime as dt
import email
import email.policy
import io
import itertools
import json
import math
import os
import posixpath
import re
import sys
import tempfile
import threading
import zipfile
from importlib.metadata import version
from pathlib import Path

PAGE_IMAGE_TOKENS = 1600
LITTLE_TEXT_CHARS_PER_PAGE = 200
SHEETS = {".xlsx", ".xls", ".ods"}
FORMATS = {
    ".docx", ".pptx", ".xlsx", ".xls", ".ods", ".odt", ".odp", ".rtf",
    ".pdf", ".html", ".htm", ".epub", ".eml", ".msg",
}


# --- estimates -------------------------------------------------------------------------


def estimate_tokens(text: str) -> int:
    """Characters / 4 for mostly-Latin text, / 3 when over 30% of letters are not Latin."""
    if not text:
        return 0
    letters = [c for c in text if c.isalpha()]
    non_latin = sum(1 for c in letters if ord(c) > 0x24F)
    divisor = 3 if letters and non_latin / len(letters) > 0.3 else 4
    return math.ceil(len(text) / divisor)


def zip_text(path: Path, patterns: list[str]) -> str:
    """The raw XML of the parts a model would have to read from an Office/ODF zip."""
    out = []
    try:
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                if any(re.fullmatch(p, name) for p in patterns):
                    out.append(z.read(name).decode("utf-8", "replace"))
    except zipfile.BadZipFile:
        return ""
    return "".join(out)


ZIP_PARTS = {
    ".docx": [r"word/document\.xml", r"word/(foot|end)notes\.xml", r"word/header\d*\.xml", r"word/footer\d*\.xml"],
    ".xlsx": [r"xl/sharedStrings\.xml", r"xl/worksheets/sheet\d+\.xml"],
    ".pptx": [r"ppt/slides/slide\d+\.xml", r"ppt/notesSlides/notesSlide\d+\.xml"],
    ".odt": [r"content\.xml"],
    ".ods": [r"content\.xml"],
    ".odp": [r"content\.xml"],
    ".epub": [r".*\.x?html?"],
}


def original_tokens(path: Path, ext: str, pages: int | None, md_text: str) -> tuple[int, bool]:
    """What reading the original would cost, and whether that is only a floor."""
    if ext == ".pdf":
        return (pages or 1) * PAGE_IMAGE_TOKENS + estimate_tokens(md_text), False
    if ext in ZIP_PARTS:
        raw = zip_text(path, ZIP_PARTS[ext])
        if raw:
            return estimate_tokens(raw), True
    if ext in {".html", ".htm", ".rtf", ".eml"}:
        return estimate_tokens(path.read_bytes().decode("utf-8", "replace")), False
    # binary formats (.xls, .msg): their bytes as text, a rough floor
    return math.ceil(path.stat().st_size / 4), True


def plural(n: int, word: str) -> str:
    return f"{n} {word}" if n == 1 else f"{n} {word}s"


def lost_items(path: Path, ext: str, extra: list[str]) -> list[str]:
    """What markdown cannot hold, counted where it can be: images, charts, objects."""
    lost: list[str] = []
    names: list[str] = []
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            if ext == ".xlsx":
                # a formula cell carries its last result in <v> - unless no program that
                # calculates ever saved the file; then the markdown shows the formula
                with_value = without = 0
                for n in names:
                    if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n):
                        continue
                    for cell in re.findall(rb"<c\b[^>]*>(.*?)</c>", z.read(n), re.DOTALL):
                        if re.search(rb"<f[ >/]", cell):
                            if re.search(rb"<v>[^<]+</v>", cell):
                                with_value += 1
                            else:
                                without += 1
                if with_value:
                    lost.append(f"{plural(with_value, 'formula')} ({'value' if with_value == 1 else 'values'} kept)")
                if without:
                    lost.append(f"{plural(without, 'formula')} with no saved result (formula shown)")
    except (zipfile.BadZipFile, OSError):
        pass
    images = [n for n in names if re.search(r"(^|/)(media|Pictures|images?)/[^/]+$", n, re.IGNORECASE)]
    charts = [n for n in names if re.search(r"/charts/chart\d+\.xml$", n)]
    objects = [n for n in names if "/embeddings/" in n or re.match(r"Object \d+/content\.xml$", n)]
    if images:
        lost.append(plural(len(images), "image"))
    if charts:
        lost.append(plural(len(charts), "chart"))
    if objects:
        lost.append(plural(len(objects), "embedded object"))
    return lost + extra


# --- markdown helpers ------------------------------------------------------------------


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ").strip()


def md_table(rows: list[list]) -> str:
    """Rows to a markdown table, the first row as the header; ragged rows padded."""
    rows = [[cell_text(v) for v in r] for r in rows]
    # trailing empty rows and columns are layout, not data
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    while width > 0 and not any(len(r) >= width and r[width - 1] for r in rows):
        width -= 1
    if width == 0:
        return ""
    rows = [(r + [""] * width)[:width] for r in rows]
    head, *body = rows
    lines = ["| " + " | ".join(head) + " |", "| " + " | ".join("---" for _ in head) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def html_to_md(html: str) -> tuple[str, int]:
    """Clean HTML to markdown: scripts, styles and navigation out, images counted."""
    from bs4 import BeautifulSoup, Comment, Declaration, Doctype, ProcessingInstruction
    from markdownify import markdownify

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "noscript", "template", "iframe", "head"]):
        tag.decompose()
    # <?xml ...?>, <!DOCTYPE ...> and comments are no text
    for node in soup.find_all(string=lambda s: isinstance(s, (Comment, Declaration, Doctype, ProcessingInstruction))):
        node.extract()
    images = 0
    for img in soup.find_all("img"):
        images += 1
        alt = img.get("alt") or ""
        img.replace_with(f"[image: {alt}]" if alt else "")
    md = markdownify(str(soup), heading_style="ATX", bullets="-")
    return re.sub(r"\n{3,}", "\n\n", md).strip() + "\n", images


# --- converters ------------------------------------------------------------------------


def docx_md(path: Path) -> tuple[str, list[str]]:
    import mammoth

    # images are left out (counted from the zip), never inlined as base64
    def no_image(image):
        return {"src": "", "alt": image.alt_text or ""}

    with open(path, "rb") as f:
        result = mammoth.convert_to_html(f, convert_image=mammoth.images.img_element(no_image))
    return html_to_md(result.value)[0], []


def pptx_md(path: Path) -> tuple[str, list[str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(str(path))
    out: list[str] = []

    def shape_md(shape) -> list[str]:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return [line for s in shape.shapes for line in shape_md(s)]
        if getattr(shape, "has_table", False):
            return [md_table([[c.text for c in row.cells] for row in shape.table.rows]), ""]
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            return [text, ""] if text else []
        return []

    for n, slide in enumerate(deck.slides, 1):
        out.append(f"<!-- Slide number: {n} -->")
        title = slide.shapes.title
        if title is not None and title.text.strip():
            out.append(f"# {title.text.strip()}")
        for shape in slide.shapes:
            if title is not None and shape.shape_id == title.shape_id:
                continue
            out.extend(shape_md(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out += ["### Notes:", notes, ""]
        out.append("")
    return "\n".join(out), []


# Spreadsheets: each reader yields (sheet name, rows, columns, row iterator), streaming, so
# a big sheet is sized and summarised without loading it whole.


def xlsx_sheets(path: Path):
    """Each cell's saved value; for a formula with no saved result (a workbook written by
    a program that does not calculate), the formula itself, so the cell is never blank."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    fwb = load_workbook(path, read_only=True, data_only=False)
    try:
        for ws, fws in zip(wb.worksheets, fwb.worksheets):
            rows, cols = ws.max_row, ws.max_column
            if rows is None or cols is None:
                ws.reset_dimensions()
                fws.reset_dimensions()
                rows = cols = 0
                for row in ws.iter_rows(values_only=True):
                    rows += 1
                    cols = max(cols, len(row))
            values = ws.iter_rows(values_only=True)
            formulas = fws.iter_rows(values_only=True)
            yield ws.title, rows or 0, cols or 0, (formula_or_value(v, f) for v, f in zip(values, formulas))
    finally:
        wb.close()
        fwb.close()


def formula_or_value(values: tuple, formulas: tuple) -> list:
    return [
        v if v is not None or not (isinstance(f, str) and f.startswith("=")) else f
        for v, f in itertools.zip_longest(values, formulas)
    ]


def xls_sheets(path: Path):
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)
    for sheet in book.sheets():
        yield sheet.name, sheet.nrows, sheet.ncols, (sheet.row_values(i) for i in range(sheet.nrows))


def qname(node) -> str:
    """An odfpy node's local name; '' for a text node, which has none."""
    name = getattr(node, "qname", None)
    return name[1] if name else ""


# An ODF table often ends in "repeat this empty row a million times": never expand that.
ODS_MAX_REPEAT = 1000


def ods_rows(table):
    from odf import table as t
    from odf import teletype

    for row in table.getElementsByType(t.TableRow):
        cells: list[str] = []
        for cell in row.childNodes:
            if qname(cell) not in ("table-cell", "covered-table-cell"):
                continue
            text = "\n".join(teletype.extractText(p) for p in cell.childNodes if qname(p) in ("p", "h"))
            repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
            cells += [text] * (min(repeat, ODS_MAX_REPEAT) if text else min(repeat, 1))
        repeat = int(row.getAttribute("numberrowsrepeated") or 1)
        for _ in range(min(repeat, ODS_MAX_REPEAT) if any(cells) else 1):
            yield cells


def ods_sheets(path: Path):
    from odf import table as t
    from odf.opendocument import load

    doc = load(str(path))
    for table in doc.spreadsheet.getElementsByType(t.Table):
        rows = list(ods_rows(table))
        while rows and not any(rows[-1]):
            rows.pop()
        cols = max((len(r) for r in rows), default=0)
        yield table.getAttribute("name"), len(rows), cols, iter(rows)


SHEET_READERS = {".xlsx": xlsx_sheets, ".xls": xls_sheets, ".ods": ods_sheets}


def sheets_md(path: Path, ext: str, max_cells: int, max_rows: int, max_tokens: int) -> tuple[str, list[str], str]:
    """The whole workbook, one table per sheet - or, past the limits, a summary."""
    sizes = [(name, rows, cols) for name, rows, cols, _ in SHEET_READERS[ext](path)]
    if all(r * c <= max_cells for _, r, c in sizes):
        out = []
        for name, _, _, rows in SHEET_READERS[ext](path):
            out += [f"## {name}", "", md_table(list(rows)), ""]
        md = "\n".join(out)
        if estimate_tokens(md) <= max_tokens:
            return md, [], "full"
    out = [
        f"# {path.name} (summary)",
        "",
        (
            "This spreadsheet is too large to read in full. Below: each sheet's size, its header "
            f"and its first {max_rows} rows. To query it, use pandas on the original: "
            f"`pd.read_excel({str(path)!r}, sheet_name=...)`."
        ),
        "",
    ]
    for name, rows, cols, it in SHEET_READERS[ext](path):
        head = []
        for row in it:
            head.append(row)
            if len(head) > max_rows:
                break
        out += [f"## Sheet: {name}", f"{rows} rows × {cols} columns", "", md_table(head), ""]
    shown = sum(min(r, max_rows + 1) for _, r, _ in sizes)
    total = sum(r for _, r, _ in sizes)
    lost = [f"rows past the first {max_rows} of each sheet ({max(total - shown, 0)} rows; summary only)"]
    return "\n".join(out), lost, "summary"


def odt_md(path: Path) -> tuple[str, list[str]]:
    """Headings, paragraphs, lists and tables, in document order: odfpy, no layout."""
    from odf import office, teletype
    from odf.opendocument import load

    doc = load(str(path))
    out: list[str] = []

    def walk(node, depth: int = 0) -> None:
        kind = qname(node)
        if kind == "h":
            level = int(node.getAttribute("outlinelevel") or 1)
            out.extend([f"{'#' * min(level, 6)} {teletype.extractText(node).strip()}", ""])
        elif kind == "p":
            text = teletype.extractText(node).strip()
            if text:
                out.extend([text, ""])
        elif kind == "list":
            for item in node.childNodes:
                if qname(item) != "list-item":
                    continue
                for child in item.childNodes:
                    name = qname(child)
                    if name == "list":
                        walk(child, depth + 1)
                    elif name in ("p", "h"):
                        text = teletype.extractText(child).strip()
                        if text:
                            out.append(f"{'  ' * depth}- {text}")
            if depth == 0:
                out.append("")
        elif kind == "table":
            out.extend([md_table(list(ods_rows(node))), ""])
        elif kind in ("section", "text", "index-body", "table-of-content"):
            for child in node.childNodes:
                walk(child, depth)

    # a loaded document's .text is a text node; the body holds the office:text element
    for root in doc.body.getElementsByType(office.Text):
        walk(root)
    return "\n".join(out), []


def odp_md(path: Path) -> tuple[str, list[str]]:
    """Text per slide, speaker notes after it: odfpy, no layout."""
    from odf import draw, presentation, teletype, text
    from odf.opendocument import load

    doc = load(str(path))
    out = []
    for i, page in enumerate(doc.getElementsByType(draw.Page), 1):
        title = page.getAttribute("name") or ""
        notes_nodes = page.getElementsByType(presentation.Notes)
        notes_text = []
        for notes in notes_nodes:
            for p in notes.getElementsByType(text.P):
                line = teletype.extractText(p).strip()
                if line:
                    notes_text.append(line)
        lines = []
        for p in page.getElementsByType(text.P) + page.getElementsByType(text.H):
            if any(_inside(p, n) for n in notes_nodes):
                continue
            line = teletype.extractText(p).strip()
            if line:
                lines.append(line)
        out.append(f"## Slide {i}" + (f": {title}" if title and not title.startswith("page") else ""))
        out.extend(lines)
        if notes_text:
            out += ["", "### Notes", *notes_text]
        out.append("")
    return "\n".join(out), []


def _inside(node, ancestor) -> bool:
    while node is not None:
        if node is ancestor:
            return True
        node = node.parentNode
    return False


def rtf_md(path: Path) -> tuple[str, list[str]]:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(path.read_bytes().decode("latin-1"), errors="ignore"), []


def pdf_md(path: Path) -> tuple[str, list[str], int]:
    """Text per page, tables as markdown tables: pdfplumber (pdfminer underneath)."""
    import pdfplumber

    out: list[str] = []
    with pdfplumber.open(path) as pdf:
        pages = len(pdf.pages)
        for n, page in enumerate(pdf.pages, 1):
            out.append(f"<!-- Page {n} -->")
            # reading order: the text above each table, the table, then what follows it
            top = page.bbox[1]
            for table in sorted(page.find_tables(), key=lambda t: t.bbox[1]):
                _, t_top, _, t_bottom = table.bbox
                if t_top > top:
                    band = page.crop((page.bbox[0], top, page.bbox[2], t_top))
                    text = (band.extract_text() or "").strip()
                    if text:
                        out.append(text)
                rows = table.extract()
                if rows:
                    out += ["", md_table(rows), ""]
                top = max(top, t_bottom)
            if top < page.bbox[3]:
                rest = page.crop((page.bbox[0], top, page.bbox[2], page.bbox[3]))
                text = (rest.extract_text() or "").strip()
                if text:
                    out.append(text)
            out.append("")
    return "\n".join(out), [], pages


def epub_md(path: Path) -> tuple[str, list[str]]:
    """The spine's chapters in reading order."""
    import warnings

    from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning

    # the package files are XML, read leniently with the HTML parser: no lxml needed
    warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)

    with zipfile.ZipFile(path) as z:
        container = BeautifulSoup(z.read("META-INF/container.xml"), "html.parser")
        opf_path = container.find("rootfile")["full-path"]
        opf = BeautifulSoup(z.read(opf_path), "html.parser")
        base = posixpath.dirname(opf_path)
        manifest = {item["id"]: item["href"] for item in opf.find_all("item") if item.get("id") and item.get("href")}
        out: list[str] = []
        for ref in opf.find_all("itemref"):
            href = manifest.get(ref.get("idref", ""))
            if not href:
                continue
            try:
                html = z.read(posixpath.normpath(posixpath.join(base, href))).decode("utf-8", "replace")
            except KeyError:
                continue
            md, _ = html_to_md(html)
            if md.strip():
                out += [md, ""]
    return "\n".join(out), []


def html_md(path: Path) -> tuple[str, list[str]]:
    md, images = html_to_md(path.read_bytes().decode("utf-8", "replace"))
    return md, [plural(images, "image")] if images else []


def mail_md(head: list[str], body: str, attachments: list[str]) -> tuple[str, list[str]]:
    lines = ["  \n".join(head), "", body.strip()]
    if attachments:
        lines += ["", "**Attachments:** " + ", ".join(attachments)]
    lost = [f"{plural(len(attachments), 'attachment')} (name kept)" if len(attachments) == 1 else f"{plural(len(attachments), 'attachment')} (names kept)"] if attachments else []
    return "\n".join(lines) + "\n", lost


def eml_md(path: Path) -> tuple[str, list[str]]:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    head = [f"**{key}:** {msg[key]}" for key in ("From", "To", "Cc", "Date", "Subject") if msg[key]]
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        content = body_part.get_content()
        body = html_to_md(content)[0] if body_part.get_content_type() == "text/html" else content
    return mail_md(head, body, [a.get_filename() or "unnamed" for a in msg.iter_attachments()])


def msg_md(path: Path) -> tuple[str, list[str]]:
    """An Outlook .msg: its MAPI properties read straight from the OLE file."""
    import olefile

    ole = olefile.OleFileIO(str(path))

    def prop(code: str, prefix: str = "") -> str:
        for kind, enc in (("001F", "utf-16-le"), ("001E", "cp1252")):
            name = f"{prefix}__substg1.0_{code}{kind}"
            if ole.exists(name):
                return ole.openstream(name).read().decode(enc, "replace").rstrip("\x00").strip()
        return ""

    try:
        head = []
        sender = prop("0C1A")
        address = prop("5D01") or prop("0C1F")
        if sender or address:
            head.append(f"**From:** {sender} <{address}>" if sender and address else f"**From:** {sender or address}")
        for label, code in (("To", "0E04"), ("Cc", "0E03"), ("Subject", "0037")):
            value = prop(code)
            if value:
                head.append(f"**{label}:** {value}")
        body = prop("1000")
        if not body and ole.exists("__substg1.0_10130102"):
            body = html_to_md(ole.openstream("__substg1.0_10130102").read().decode("utf-8", "replace"))[0]
        attachments = []
        for entry in ole.listdir(streams=False, storages=True):
            if len(entry) == 1 and entry[0].startswith("__attach_version1.0_"):
                prefix = entry[0] + "/"
                attachments.append(prop("3707", prefix) or prop("3704", prefix) or "unnamed")
    finally:
        ole.close()
    return mail_md(head, body, attachments)


CONVERTERS = {
    ".docx": (docx_md, "mammoth"),
    ".pptx": (pptx_md, "python-pptx"),
    ".odt": (odt_md, "odfpy"),
    ".odp": (odp_md, "odfpy"),
    ".rtf": (rtf_md, "striprtf"),
    ".epub": (epub_md, "markdownify"),
    ".html": (html_md, "markdownify"),
    ".htm": (html_md, "markdownify"),
    ".eml": (eml_md, None),
    ".msg": (msg_md, "olefile"),
}
SHEET_LIBS = {".xlsx": "openpyxl", ".xls": "xlrd", ".ods": "odfpy"}


def converter_name(ext: str) -> str:
    if ext == ".pdf":
        return f"pdfplumber {version('pdfplumber')}"
    if ext in SHEET_LIBS:
        return f"{SHEET_LIBS[ext]} {version(SHEET_LIBS[ext])}"
    lib = CONVERTERS[ext][1]
    return f"{lib} {version(lib)}" if lib else "python email"


def versions() -> dict:
    libs = ["mammoth", "markdownify", "python-pptx", "openpyxl", "xlrd", "odfpy", "pdfplumber", "striprtf", "olefile"]
    return {"python": sys.version.split()[0], **{lib: version(lib) for lib in libs}}


def convert(src: Path, out: Path, max_cells: int, max_rows: int, max_tokens: int) -> dict:
    ext = src.suffix.lower()
    if ext not in FORMATS:
        raise SystemExit(f"not a format VeroDotto converts: {ext}")
    kind = "full"
    pages = None
    keep = None

    if ext in SHEETS:
        md, extra_lost, kind = sheets_md(src, ext, max_cells, max_rows, max_tokens)
    elif ext == ".pdf":
        md, extra_lost, pages = pdf_md(src)
        text_chars = len(re.sub(r"<!-- Page \d+ -->|\s", "", md))
        if text_chars < LITTLE_TEXT_CHARS_PER_PAGE * max(pages, 1):
            keep = "little-text"
    else:
        md, extra_lost = CONVERTERS[ext][0](src)

    md = re.sub(r"\n{3,}", "\n\n", md.replace("\r\n", "\n").replace("\r", "\n")).strip() + "\n"
    if not md.strip() and keep is None:
        keep = "empty"
    orig, is_floor = original_tokens(src, ext, pages, md)
    meta = {
        "source": src.name,
        "ext": ext,
        "size": src.stat().st_size,
        "mdChars": len(md),
        "mdTokens": estimate_tokens(md),
        "origTokens": orig,
        "isLowerBound": is_floor,
        "lost": lost_items(src, ext, extra_lost),
        "converter": converter_name(ext),
        "kind": kind,
        "convertedAt": dt.datetime.now(dt.UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }
    if keep:
        meta["keep"] = keep
    if pages is not None:
        meta["pages"] = pages
    out.mkdir(parents=True, exist_ok=True)
    (out / "doc.md").write_text(md, encoding="utf-8", newline="\n")
    (out / "meta.json").write_text(json.dumps(meta, ensure_ascii=False) + "\n", encoding="utf-8")
    return meta


# --- self-test -------------------------------------------------------------------------

WORD = "Selftest"


def tiny_pdf() -> bytes:
    """A one-page PDF saying the word, xref offsets computed: no library writes PDFs here."""
    line = f"{WORD} page with enough words to read as text, " * 6
    stream = f"BT /F1 9 Tf 20 720 Td ({line}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def tiny_docx(dest: Path) -> None:
    with zipfile.ZipFile(dest, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="r1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
        )
        z.writestr(
            "word/document.xml",
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{WORD} hello</w:t></w:r></w:p></w:body></w:document>",
        )


def tiny_epub(dest: Path) -> None:
    with zipfile.ZipFile(dest, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        z.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Book</dc:title></metadata>'
            '<manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/></spine></package>',
        )
        z.writestr("OEBPS/c1.xhtml", f"<html><body><h1>{WORD}</h1><p>hello</p></body></html>")


def selftest() -> None:
    """One tiny sample per format, built here, so setup proves each converter runs."""
    from odf import draw, style, table, text
    from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText
    from openpyxl import Workbook
    from pptx import Presentation

    with tempfile.TemporaryDirectory() as tmp:
        t = Path(tmp)
        tiny_docx(t / "a.docx")
        tiny_epub(t / "a.epub")
        (t / "a.pdf").write_bytes(tiny_pdf())
        (t / "a.rtf").write_text(r"{\rtf1\ansi " + WORD + r" hello\par}", encoding="ascii")
        (t / "a.html").write_text(f"<html><body><nav>menu</nav><h1>{WORD}</h1></body></html>", encoding="utf-8")
        (t / "a.eml").write_text(f"From: a@b.c\nSubject: {WORD}\n\nhello\n", encoding="utf-8")
        wb = Workbook()
        wb.active.append([WORD, 1])
        wb.save(t / "a.xlsx")
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = WORD
        deck.save(t / "a.pptx")
        odt = OpenDocumentText()
        odt.text.addElement(text.H(outlinelevel=1, text=WORD))
        odt.save(str(t / "a.odt"))
        ods = OpenDocumentSpreadsheet()
        sheet = table.Table(name="S")
        row = table.TableRow()
        cell = table.TableCell(valuetype="string")
        cell.addElement(text.P(text=WORD))
        row.addElement(cell)
        sheet.addElement(row)
        ods.spreadsheet.addElement(sheet)
        ods.save(str(t / "a.ods"))
        odp = OpenDocumentPresentation()
        layout = style.PageLayout(name="L")
        odp.automaticstyles.addElement(layout)
        master = style.MasterPage(name="M", pagelayoutname=layout)
        odp.masterstyles.addElement(master)
        page = draw.Page(masterpagename=master, name="One")
        frame = draw.Frame(width="10cm", height="2cm", x="1cm", y="1cm")
        box = draw.TextBox()
        box.addElement(text.P(text=WORD))
        frame.addElement(box)
        page.addElement(frame)
        odp.presentation.addElement(page)
        odp.save(str(t / "a.odp"))

        for name in sorted(p.name for p in t.iterdir()):
            meta = convert(t / name, t / f"out-{name}", 5000, 20, 50000)
            md = (t / f"out-{name}" / "doc.md").read_text(encoding="utf-8")
            if WORD not in md:
                raise SystemExit(f"selftest: {name} did not convert ({meta['converter']})")
        # .xls and .msg cannot be written here; their readers are imported at least
        import olefile  # noqa: F401
        import xlrd  # noqa: F401
    print(json.dumps(versions()))


# --- entry -----------------------------------------------------------------------------


def main() -> None:
    if sys.argv[1:] == ["--selftest"]:
        selftest()
        return
    # setup's step 2: uv has installed the dependencies by the time this runs
    if sys.argv[1:] == ["--deps"]:
        return
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["convert"])
    parser.add_argument("src")
    parser.add_argument("out")
    parser.add_argument("--max-sheet-cells", type=int, default=5000)
    parser.add_argument("--max-sheet-rows", type=int, default=20)
    parser.add_argument("--max-tokens", type=int, default=50000)
    parser.add_argument("--timeout", type=float, default=100)
    args = parser.parse_args()
    # A conversion that hangs (a hostile PDF) must not hold the read: past the limit the
    # process ends and the plugin reads the original. A timer works on every OS.
    watchdog = threading.Timer(args.timeout, lambda: os._exit(124))
    watchdog.daemon = True
    watchdog.start()
    convert(Path(args.src), Path(args.out), args.max_sheet_cells, args.max_sheet_rows, args.max_tokens)
    watchdog.cancel()


if __name__ == "__main__":
    main()
